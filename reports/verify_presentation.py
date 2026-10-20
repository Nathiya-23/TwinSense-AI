"""
Verification script for TwinSense AI PPTX and PDF presentations.
"""
from pathlib import Path
import pptx
import pypdf

def main():
    pptx_path = Path("reports/TwinSense_AI_Presentation.pptx")
    pdf_path = Path("reports/TwinSense_AI_Presentation.pdf")
    
    assert pptx_path.exists(), "PPTX file does not exist!"
    assert pdf_path.exists(), "PDF file does not exist!"
    
    # 1. PPTX Check
    prs = pptx.Presentation(str(pptx_path))
    print(f"PPTX verification:")
    print(f" - Path: {pptx_path}")
    print(f" - Size: {pptx_path.stat().st_size} bytes")
    print(f" - Number of slides: {len(prs.slides)}")
    assert len(prs.slides) == 12, f"Expected 12 slides, got {len(prs.slides)}"
    
    # 2. PDF Check
    reader = pypdf.PdfReader(str(pdf_path))
    print(f"\nPDF verification:")
    print(f" - Path: {pdf_path}")
    print(f" - Size: {pdf_path.stat().st_size} bytes")
    print(f" - Number of pages: {len(reader.pages)}")
    assert len(reader.pages) == 12, f"Expected 12 pages, got {len(reader.pages)}"
    
    # Extract text from all pages
    full_text = ""
    for i, page in enumerate(reader.pages):
        t = page.extract_text()
        full_text += f"\n--- Page {i+1} ---\n{t}"
        
    disclosures = [
        "Synthetic Data",
        "89.29%",
        "Hypoglycemia",
        "Not supported",
        "Not Clinically Validated",
        "TwinSense AI",
        "Problem Statement",
        "System Architecture",
        "Random Forest",
        "What-If"
    ]
    
    print("\nRequired disclosures & keywords check in PDF:")
    for d in disclosures:
        found = d.lower() in full_text.lower()
        status = "FOUND" if found else "MISSING"
        print(f" - '{d}': {status}")
        assert found, f"Required content '{d}' missing from PDF text!"
        
    print("\nSUCCESS: Both PPTX and PDF files verified successfully!")

if __name__ == "__main__":
    main()
